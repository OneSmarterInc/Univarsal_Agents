"""Bounded local parsing. CLI child emits only structured data to its parent."""

import io
import json
import os
import subprocess
import sys
import tempfile
import warnings
import zipfile
from pathlib import Path

from PIL import Image, ImageOps

from .errors import ServiceError

MAX_BYTES = 10 * 1024 * 1024
MAX_PIXELS = 25_000_000
MAX_TEXT = 36_000
Image.MAX_IMAGE_PIXELS = MAX_PIXELS


def docx_guard(raw):
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            entries = z.infolist()
            if (
                len(entries) > 2000
                or sum(i.file_size for i in entries) > 50 * 1024 * 1024
                or any(i.file_size > max(1, i.compress_size) * 100 for i in entries)
            ):
                raise ServiceError("decompression_limit")
            names = {i.filename for i in entries}
            if "word/document.xml" not in names or "[Content_Types].xml" not in names:
                raise ServiceError("invalid_document")
            for i in entries:
                if (
                    i.flag_bits & 1
                    or i.filename.startswith("/")
                    or ".." in Path(i.filename).parts
                    or "vbaproject" in i.filename.lower()
                ):
                    raise ServiceError("invalid_document")
                if i.filename.endswith(".rels"):
                    # External hyperlinks are inert; other external relationships are rejected.
                    from defusedxml.ElementTree import fromstring

                    for relation in fromstring(z.read(i)):
                        if relation.attrib.get("TargetMode") == "External" and not relation.attrib.get(
                            "Type", ""
                        ).endswith("/hyperlink"):
                            raise ServiceError("external_document_relationship")
    except ServiceError:
        raise
    except Exception as exc:
        raise ServiceError("invalid_document") from exc


def verify_upload(raw, filename, media_type, kind):
    if not raw:
        raise ServiceError("empty_file")
    if len(raw) > MAX_BYTES:
        raise ServiceError("file_too_large", 413)
    suffix = Path(filename or "").suffix.lower()
    if kind == "cv" and suffix == ".pdf" and media_type == "application/pdf":
        if not raw.startswith(b"%PDF-"):
            raise ServiceError("invalid_document")
    elif (
        kind == "cv"
        and suffix == ".docx"
        and media_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    ):
        docx_guard(raw)
    elif kind == "linkedin" and suffix in {".png", ".jpg", ".jpeg"}:
        expected = "PNG" if suffix == ".png" else "JPEG"
        if media_type != ("image/png" if expected == "PNG" else "image/jpeg"):
            raise ServiceError("unsupported_media_type", 415)
        try:
            with warnings.catch_warnings():
                warnings.simplefilter("error", Image.DecompressionBombWarning)
                with Image.open(io.BytesIO(raw)) as image:
                    if image.format != expected:
                        raise ServiceError("unsupported_media_type", 415)
                    if image.width * image.height > MAX_PIXELS or max(image.size) > 20_000:
                        raise ServiceError("image_pixel_limit")
                    image.verify()
        except ServiceError:
            raise
        except Exception as exc:
            raise ServiceError("invalid_image") from exc
    else:
        raise ServiceError("unsupported_media_type", 415)
    return suffix


def ocr(path):
    try:
        with Image.open(path) as image:
            if image.width * image.height > MAX_PIXELS:
                raise ServiceError("image_pixel_limit")
            image = ImageOps.exif_transpose(image).convert("RGB")
            with tempfile.TemporaryDirectory(prefix="ocr-") as directory:
                normalized = Path(directory) / "image.png"
                image.save(normalized)
                texts = []
                for mode in (3, 11):
                    proc = subprocess.run(
                        ["tesseract", str(normalized), "stdout", "-l", "eng", "--psm", str(mode)],
                        capture_output=True,
                        timeout=60,
                        check=True,
                    )
                    text = proc.stdout.decode("utf-8", errors="replace").strip()
                    texts.append(text)
                    if len(text) >= 60:
                        break
                return max(texts, key=len)
    except ServiceError:
        raise
    except (OSError, subprocess.SubprocessError) as exc:
        raise ServiceError("ocr_failed", retryable=True) from exc


def parse_file(path, kind):
    path = Path(path)
    blocks, notices = [], []
    page_count = None
    ocr_pages = 0

    def add(value, locator, origin):
        value = value.strip()
        if value:
            blocks.append({"text": value, "locator": locator, "origin": origin})
        if sum(len(b["text"]) for b in blocks) > MAX_TEXT:
            raise ServiceError("text_limit_exceeded")

    if kind == "linkedin":
        add(ocr(path), {"bounding_box": None}, "ocr")
    elif path.suffix == ".docx":
        from docx import Document

        docx_guard(path.read_bytes())
        document = Document(path)
        for i, p in enumerate(document.paragraphs):
            add(p.text, {"paragraph": i + 1}, "native_text")
        for ti, table in enumerate(document.tables):
            for ri, row in enumerate(table.rows):
                for ci, cell in enumerate(row.cells):
                    add(cell.text, {"table": ti + 1, "row": ri + 1, "cell": ci + 1}, "native_text")
        if document.inline_shapes or any(p.text for s in document.sections for p in s.header.paragraphs):
            notices.append("docx_headers_drawings_and_embedded_images_not_extracted")
    else:
        from pypdf import PdfReader

        reader = PdfReader(path, strict=True)
        if reader.is_encrypted:
            raise ServiceError("encrypted_document_unsupported")
        page_count = len(reader.pages)
        if page_count > 30:
            raise ServiceError("page_limit_exceeded")
        for i, page in enumerate(reader.pages):
            native = page.extract_text() or ""
            add(native, {"page": i + 1}, "native_text")
            # Image-bearing pages also undergo OCR: native text alone can miss scanned regions.
            resources = page.get("/Resources", {})
            has_images = bool(resources.get("/XObject")) if hasattr(resources, "get") else True
            if len(native.strip()) < 60 or has_images:
                with tempfile.TemporaryDirectory(prefix="pdf-") as directory:
                    target = Path(directory) / "page"
                    try:
                        subprocess.run(
                            [
                                "pdftoppm",
                                "-f",
                                str(i + 1),
                                "-l",
                                str(i + 1),
                                "-singlefile",
                                "-scale-to",
                                "2200",
                                "-png",
                                str(path),
                                str(target),
                            ],
                            capture_output=True,
                            check=True,
                            timeout=60,
                        )
                        text = ocr(target.with_suffix(".png"))
                    except (OSError, subprocess.SubprocessError) as exc:
                        raise ServiceError("pdf_render_failed") from exc
                ocr_pages += 1
                if " ".join(text.split()) != " ".join(native.split()):
                    add(text, {"page": i + 1, "bounding_box": None}, "ocr")
    if not blocks:
        raise ServiceError("unreadable_document")
    return {
        "blocks": blocks,
        "warnings": notices,
        "coverage": {
            "pages_total": page_count,
            "pages_read": page_count,
            "pages_ocr": ocr_pages,
            "characters_read": sum(len(b["text"]) for b in blocks),
        },
    }


def parse_isolated(path, kind, timeout=180):
    # No API keys, OAuth tokens or application secrets in the parser environment.
    env = {k: v for k, v in os.environ.items() if k in {"PATH", "LANG", "LC_ALL", "TESSDATA_PREFIX"}}
    try:
        import signal

        with tempfile.TemporaryDirectory(prefix="agents-parser-") as directory:
            env["TMPDIR"] = directory
            proc = subprocess.Popen(
                [sys.executable, "-m", "extraction_service.parsers", str(path), kind],
                env=env,
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
                start_new_session=True,
            )
            try:
                stdout, _ = proc.communicate(timeout=timeout)
            except subprocess.TimeoutExpired:
                os.killpg(proc.pid, signal.SIGKILL)
                proc.communicate()
                raise
            if len(stdout) > 2 * 1024 * 1024:
                raise ServiceError("parser_output_limit")
            data = json.loads(stdout)
        if "error" in data:
            raise ServiceError(data["error"])
        if proc.returncode:
            raise ServiceError("invalid_document")
        return data
    except ServiceError:
        raise
    except subprocess.TimeoutExpired as exc:
        raise ServiceError("parser_timeout", retryable=True) from exc
    except Exception as exc:
        raise ServiceError("invalid_document") from exc


def main():
    import resource

    resource.setrlimit(resource.RLIMIT_AS, (1536 * 1024**2, 1536 * 1024**2))
    resource.setrlimit(resource.RLIMIT_CPU, (120, 120))
    resource.setrlimit(resource.RLIMIT_FSIZE, (100 * 1024**2, 100 * 1024**2))
    resource.setrlimit(resource.RLIMIT_NOFILE, (128, 128))
    try:
        result = parse_file(sys.argv[1], sys.argv[2])
    except ServiceError as exc:
        result = {"error": exc.code}
    except Exception:  # noqa: BLE001 - parser boundary must not expose document content
        result = {"error": "invalid_document"}
    print(json.dumps(result))


if __name__ == "__main__":
    main()

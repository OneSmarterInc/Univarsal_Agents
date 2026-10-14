import io
import shutil
import zipfile
from datetime import date

import pytest
from docx import Document
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

from extraction_service.errors import ServiceError
from extraction_service.extractors import chunks, dates_sane, extract_blocks
from extraction_service.parsers import docx_guard, parse_file, parse_isolated, verify_upload
from extraction_service.schemas import Claim, Observations


def test_pdf_and_docx_parsing(tmp_path, pdf):
    path = tmp_path / "cv.pdf"
    path.write_bytes(pdf)
    assert "Alex Example" in parse_isolated(path, "cv")["blocks"][0]["text"]
    document = Document()
    document.add_paragraph("Alex Example")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Python"
    path = tmp_path / "cv.docx"
    document.save(path)
    parsed = parse_isolated(path, "cv")
    assert any(b["text"] == "Python" and "table" in b["locator"] for b in parsed["blocks"])


def test_pdf_page_limit(tmp_path):
    path = tmp_path / "long.pdf"
    c = canvas.Canvas(str(path))
    for _ in range(31):
        c.drawString(50, 700, "Synthetic page")
        c.showPage()
    c.save()
    with pytest.raises(ServiceError, match="page_limit_exceeded"):
        parse_file(path, "cv")


def test_text_limit(tmp_path):
    path = tmp_path / "long.docx"
    doc = Document()
    doc.add_paragraph("x" * 36001)
    doc.save(path)
    with pytest.raises(ServiceError, match="text_limit_exceeded"):
        parse_file(path, "cv")


@pytest.mark.parametrize(
    "name,mime,kind,raw",
    [
        ("x.pdf", "application/pdf", "cv", b"not PDF"),
        ("x.txt", "text/plain", "cv", b"Alex Example"),
        ("x.png", "image/png", "linkedin", b"not PNG"),
        (
            "x.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "cv",
            b"not ZIP",
        ),
    ],
)
def test_invalid_uploads(name, mime, kind, raw):
    with pytest.raises(ServiceError):
        verify_upload(raw, name, mime, kind)


def test_file_size_limit():
    with pytest.raises(ServiceError, match="file_too_large"):
        verify_upload(b"%PDF-" + b"x" * (10 * 1024**2), "cv.pdf", "application/pdf", "cv")


def test_zip_bomb():
    data = io.BytesIO()
    with zipfile.ZipFile(data, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("word/document.xml", b"x" * 100000)
        z.writestr("[Content_Types].xml", "x")
    with pytest.raises(ServiceError, match="decompression_limit"):
        docx_guard(data.getvalue())


def test_external_relationship_rejected():
    data = io.BytesIO()
    with zipfile.ZipFile(data, "w") as z:
        z.writestr("word/document.xml", "x")
        z.writestr("[Content_Types].xml", "x")
        z.writestr(
            "word/_rels/document.xml.rels",
            '<Relationships><Relationship TargetMode="External" Type="image" Target="http://169.254.169.254/"/></Relationships>',
        )
    with pytest.raises(ServiceError, match="external_document_relationship"):
        docx_guard(data.getvalue())


def test_real_scanned_pdf_ocr(tmp_path):
    if not shutil.which("tesseract") or not shutil.which("pdftoppm"):
        pytest.skip("Local OCR binaries unavailable")
    image = Image.new("RGB", (1400, 500), "white")
    font = ImageFont.truetype("DejaVuSans.ttf", 55)
    ImageDraw.Draw(image).text((70, 70), "Alex Example\nPython Developer", fill="black", font=font)
    path = tmp_path / "scanned.pdf"
    c = canvas.Canvas(str(path), pagesize=(700, 250))
    c.drawImage(ImageReader(image), 0, 0, width=700, height=250)
    c.save()
    parsed = parse_isolated(path, "cv")
    assert parsed["coverage"]["pages_ocr"] == 1
    assert any("Alex Example" in b["text"] for b in parsed["blocks"])


def test_date_validation_has_clock_and_reversed_dates():
    assert dates_sane("2024", "2025", date(2026, 1, 1))
    assert dates_sane("Jan 2024", "Present", date(2026, 1, 1))
    assert not dates_sane("2025", "2024", date(2026, 1, 1))
    assert not dates_sane("2090", None, date(2026, 1, 1))
    assert not dates_sane("tomorrow", None, date(2026, 1, 1))


def test_chunks_cover_unicode_without_discarding_suffix():
    text = "测试" * 5000 + "FINAL_MARKER"
    pieces = list(chunks(text))
    assert all(len(p.encode()) <= 2200 for _, p in pieces)
    assert pieces[-1][1].endswith("FINAL_MARKER")
    covered = set()
    for start, piece in pieces:
        covered.update(range(start, start + len(piece)))
    assert covered == set(range(len(text)))


def test_unsupported_claim_and_contact_fields_are_not_fabricated():
    class Fake:
        def extract(self, kind, block, check):
            return Observations(
                scalars=[
                    Claim(field="name", value="Invented Person", quote="Alex Example", block_id="b"),
                    Claim(field="email", value="a@example.test", quote="a@example.test", block_id="b"),
                ]
            )

    block = {
        "id": "b",
        "input_id": "i",
        "text": "Alex Example a@example.test",
        "origin": "ocr",
        "locator": {},
    }
    result = extract_blocks("linkedin", [block], Fake(), requested=["name", "email"])
    assert result["data"] == {"name": None, "email": None}
    assert result["fields"]["/email"]["status"] == "unsupported"
    assert result["warnings"]


def test_conflicts_retain_evidence_and_do_not_pick_a_person():
    class Fake:
        def extract(self, kind, block, check):
            return Observations(
                scalars=[Claim(field="name", value=block["text"], quote=block["text"], block_id=block["id"])]
            )

    blocks = [
        {"id": str(i), "input_id": str(i), "text": name, "origin": "ocr", "locator": {}}
        for i, name in enumerate(["Alex Example", "Sam Example"])
    ]
    result = extract_blocks("linkedin", blocks, Fake())
    assert result["data"]["name"] is None
    assert result["fields"]["/name"]["status"] == "unavailable"
    assert len(result["fields"]["/name"]["evidence_ids"]) == 2
